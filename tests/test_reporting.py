import csv
import hashlib
import json
from dataclasses import replace
from zipfile import ZipFile

import pytest
from conftest import HEADERS, VALID
from openpyxl import load_workbook

from sheetmerge import (
    Config,
    ExportError,
    ReconciliationError,
    consolidate,
    export_batch,
    load_batch,
    reporting,
)


def read_csv(path):
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def test_master_exports_preserve_exact_values_and_literal_excel_text(
    tmp_path, write_csv, config_factory
):
    inputs = tmp_path / "inputs"
    path = write_csv(
        inputs / "sales.csv",
        rows=[
            ["002", "0001", "2024-02-29", "0007", "-12.3400", "USD"],
            [
                "001",
                "=1+1",
                "2024-02-29",
                "customer\r\nnext",
                "1234567890123456789012345.67",
                "USD",
            ],
            ["001", "+0002", "2024-02-29", "@literal", "10", "JPY"],
            ["001", "0003", "2024-02-29", "0007", "0.001", "KWD"],
        ],
    )
    before = path.read_bytes(), path.stat().st_mtime_ns
    config = config_factory()
    result = consolidate(load_batch(inputs, config))
    output = tmp_path / "out"
    exported = export_batch(result, config, output)
    assert exported.output_directory == output
    assert exported.reconciliation.counts == (4, 4, 0, 0)
    rows = read_csv(output / "master.csv")
    assert list(rows[0]) == [*HEADERS, "source_file", "source_sheet", "source_row"]
    assert [(row["branch_id"], row["transaction_id"]) for row in rows] == [
        ("001", "+0002"),
        ("001", "0003"),
        ("001", "=1+1"),
        ("002", "0001"),
    ]
    assert [row["amount"] for row in rows] == [
        "10",
        "0.001",
        "1234567890123456789012345.67",
        "-12.34",
    ]
    workbook = load_workbook(output / "master.xlsx", data_only=False)
    try:
        actual = []
        for cells in workbook["Master"].iter_rows(min_row=2):
            assert all(cell.data_type == "s" for cell in cells if cell.value is not None)
            actual.append(["" if cell.value is None else cell.value for cell in cells])
        assert actual == [list(row.values()) for row in rows]
        assert workbook["Master"].freeze_panes == "A2"
    finally:
        workbook.close()
    assert (path.read_bytes(), path.stat().st_mtime_ns) == before
    assert list(inputs.iterdir()) == [path]


def test_exclusion_audit_and_reconciliation_reports_cover_every_record(
    tmp_path, write_csv, config_factory
):
    inputs = tmp_path / "inputs"
    write_csv(
        inputs / "sales.csv",
        rows=[
            VALID,
            VALID,
            ["001", "conflict", VALID[2], "C", "1.00", "USD"],
            ["001", "conflict", VALID[2], "C", "2.00", "USD"],
            ["001", "invalid", "bad", "C", "3.00", "USD"],
            ["001", "bad-money", VALID[2], "C", "invalid", "USD"],
            [""] * 6,
        ],
    )
    config = config_factory()
    export_batch(consolidate(load_batch(inputs, config)), config, tmp_path / "out")
    output = tmp_path / "out"
    excluded = read_csv(output / "exclusions.csv")
    assert len(excluded) == 6
    assert {row["status"] for row in excluded} == {"rejected", "redundant"}
    codes = [code for row in excluded for code in json.loads(row["reason_codes"])]
    assert codes.count("CONFLICTING_TRANSACTION") == 2
    assert {"EXACT_DUPLICATE", "INVALID_DATE", "INVALID_AMOUNT", "BLANK_ROW"} <= set(codes)
    duplicate = next(row for row in excluded if row["status"] == "redundant")
    assert duplicate["duplicate_of_file"] == "sales.csv"
    assert duplicate["duplicate_of_row"] == "2"
    invalid = next(row for row in excluded if "INVALID_DATE" in row["reason_codes"])
    assert json.loads(invalid["raw_values"])["transaction_date"] == "bad"
    audit = [json.loads(line) for line in (output / "audit.jsonl").read_text().splitlines()]
    assert len(audit) == 7
    assert {row["source"]["source_row"] for row in audit} == set(range(2, 9))
    assert sum(row["status"] == "accepted" for row in audit) == 1
    report = json.loads((output / "reconciliation.json").read_text())
    assert report["counts"] == {"input": 7, "accepted": 1, "rejected": 5, "redundant": 1}
    assert report["unquantified_count"] == 2
    assert report["monetary_complete"] is False
    usd = next(row for row in read_csv(output / "reconciliation.csv") if row["currency"] == "USD")
    assert usd["input_known_amount"] == "30.68"
    assert usd["accepted_known_amount"] == usd["redundant_known_amount"] == "12.34"
    assert usd["rejected_known_amount"] == "6.00"
    assert usd["monetary_complete"] == "false"
    assert usd["input_unquantified_count"] == "1"


def test_all_artifacts_are_deterministic_and_manifest_hashes_match(
    tmp_path, write_csv, config_factory
):
    inputs = tmp_path / "inputs"
    write_csv(inputs / "z.csv", rows=[VALID, VALID])
    write_csv(inputs / "a.csv", rows=[["002", *VALID[1:]]])
    config = config_factory()
    batch = load_batch(inputs, config)
    output_a, output_b = tmp_path / "a", tmp_path / "b"
    export_batch(consolidate(batch), config, output_a)
    reversed_batch = replace(
        batch, rows=tuple(reversed(batch.rows)), files=tuple(reversed(batch.files))
    )
    export_batch(consolidate(reversed_batch), config, output_b)
    assert {path.name: path.read_bytes() for path in output_a.iterdir()} == {
        path.name: path.read_bytes() for path in output_b.iterdir()
    }
    manifest = json.loads((output_a / "manifest.json").read_text())
    assert manifest["business_key"] == ["branch_id", "transaction_id"]
    assert manifest["configuration"]["currency_precision"]["USD"] == 2
    assert manifest["configuration_sha256"] == batch.config_fingerprint == config.fingerprint
    for name, digest in manifest["output_sha256"].items():
        assert hashlib.sha256((output_a / name).read_bytes()).hexdigest() == digest
    with ZipFile(output_a / "master.xlsx") as archive:
        assert all(entry.date_time == (2000, 1, 1, 0, 0, 0) for entry in archive.infolist())
        assert archive.read("docProps/core.xml").count(b"2000-01-01T00:00:00Z") == 2


@pytest.fixture
def exportable(tmp_path, write_csv, config_factory):
    inputs = tmp_path / "inputs"
    write_csv(inputs / "sales.csv")
    config = config_factory()
    return consolidate(load_batch(inputs, config)), config


def test_existing_outputs_are_never_overwritten(tmp_path, exportable):
    output = tmp_path / "out"
    output.mkdir()
    sentinel = output / "master.csv"
    sentinel.write_text("existing contents")
    with pytest.raises(ExportError, match="already exists"):
        export_batch(*exportable, output)
    assert sentinel.read_text() == "existing contents"


def test_output_inside_input_tree_is_rejected(exportable):
    result, config = exportable
    output = result.batch.input_root / "out"
    with pytest.raises(ExportError, match="outside the input tree"):
        export_batch(result, config, output)
    assert not output.exists()


@pytest.mark.parametrize(
    "value",
    ["a" * 32768, "illegal\x01text", "illegal\ufffftext"],
    ids=["overlong", "control-character", "xml-noncharacter"],
)
def test_excel_limits_fail_instead_of_truncating(tmp_path, write_csv, config_factory, value):
    inputs = tmp_path / "inputs"
    write_csv(inputs / "sales.csv", rows=[[value, *VALID[1:]]])
    config = config_factory()
    with pytest.raises(ExportError, match="Excel"):
        export_batch(consolidate(load_batch(inputs, config)), config, tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_excel_row_limit_is_checked_before_writing(tmp_path, exportable, monkeypatch):
    monkeypatch.setattr(reporting, "EXCEL_MAX_ROWS", 1)
    with pytest.raises(ExportError, match="row limit"):
        export_batch(*exportable, tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_writer_failure_leaves_no_partial_bundle(tmp_path, exportable, monkeypatch):
    def fail(*args):
        raise OSError("disk full")

    monkeypatch.setattr(reporting, "_write_workbook", fail)
    with pytest.raises(ExportError, match="disk full"):
        export_batch(*exportable, tmp_path / "out")
    assert not (tmp_path / "out").exists()
    assert not list(tmp_path.glob(".out-*"))


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_readback_detects_incorrect_master_and_prevents_publication(
    tmp_path, exportable, monkeypatch, format
):
    if format == "csv":
        original = reporting._write_csv

        def corrupt_csv(path, headers, rows):
            if path.name == "master.csv":
                rows = [row.copy() for row in rows]
                rows[0][4] = "99.00"
            original(path, headers, rows)

        monkeypatch.setattr(reporting, "_write_csv", corrupt_csv)
    else:
        original = reporting._write_workbook

        def corrupt_excel(path, rows):
            original(path, rows)
            workbook = load_workbook(path)
            workbook["Master"]["A2"] = "=1+1"
            workbook.save(path)
            workbook.close()

        monkeypatch.setattr(reporting, "_write_workbook", corrupt_excel)
    with pytest.raises(ExportError, match="verification|formula"):
        export_batch(*exportable, tmp_path / "out")
    assert not (tmp_path / "out").exists()
    assert not list(tmp_path.glob(".out-*"))


def test_reconciliation_failure_writes_nothing(tmp_path, exportable, monkeypatch):
    def fail(*args):
        raise ReconciliationError("ledger mismatch")

    monkeypatch.setattr(reporting, "reconcile", fail)
    with pytest.raises(ReconciliationError, match="ledger mismatch"):
        export_batch(*exportable, tmp_path / "new-parent" / "out")
    assert not (tmp_path / "new-parent").exists()


def test_export_rechecks_a_lost_record_instead_of_trusting_prior_state(tmp_path, exportable):
    result, config = exportable
    with pytest.raises(ReconciliationError):
        export_batch(replace(result, rows=()), config, tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_header_only_batch_exports_headers_and_zero_counts(tmp_path, write_csv, config_factory):
    inputs = tmp_path / "inputs"
    write_csv(inputs / "empty.csv", rows=[])
    config = config_factory()
    result = export_batch(consolidate(load_batch(inputs, config)), config, tmp_path / "out")
    assert result.reconciliation.counts == (0, 0, 0, 0)
    assert read_csv(tmp_path / "out" / "master.csv") == []
    assert read_csv(tmp_path / "out" / "exclusions.csv") == []


@pytest.mark.parametrize(
    "change",
    [
        "precision",
        "normalization",
        "mapping",
        "date_format",
        "profile_name",
        "patterns",
    ],
)
def test_export_rejects_configuration_not_applied_at_ingestion(tmp_path, exportable, change):
    result, original = exportable
    profile = original.profiles[0]
    if change == "precision":
        config = replace(original, currency_precision={**original.currency_precision, "USD": 3})
    else:
        overrides = {
            "normalization": {"trim_whitespace": True},
            "mapping": {"columns": {**profile.columns, "amount": "Total"}},
            "date_format": {"date_formats": ("%d/%m/%Y",)},
            "profile_name": {"name": "changed"},
            "patterns": {"patterns": ("other*.csv",)},
        }
        config = replace(original, profiles=(replace(profile, **overrides[change]),))
    target = tmp_path / "not-created" / "out"
    with pytest.raises(ReconciliationError, match="Configuration mismatch"):
        export_batch(result, config, target)
    assert not target.parent.exists()


def test_equivalent_reloaded_configuration_is_allowed(tmp_path, exportable):
    result, config = exportable
    profile = config.profiles[0]
    equivalent = Config(
        dict(reversed(list(config.currency_precision.items()))),
        (replace(profile, columns=dict(reversed(list(profile.columns.items())))),),
    )
    assert config is not equivalent
    output = export_batch(result, equivalent, tmp_path / "out")
    assert output.reconciliation.counts == (1, 1, 0, 0)
