import hashlib
from datetime import datetime
from decimal import Decimal

import pytest
from conftest import HEADERS, VALID
from openpyxl import load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from sheetmerge import ConfigurationError, IngestionError, discover_inputs, load_batch


def test_recursive_discovery_is_sorted_and_filters_formats(tmp_path):
    for name in ("z.xlsx", "nested/a.CSV", "a.csv", "ignored.txt", "old.xls"):
        path = tmp_path / name
        path.parent.mkdir(exist_ok=True)
        path.touch()
    assert [path.relative_to(tmp_path).as_posix() for path in discover_inputs(tmp_path)] == [
        "a.csv",
        "nested/a.CSV",
        "z.xlsx",
    ]


def test_missing_root_and_empty_batch(tmp_path, config_factory):
    with pytest.raises(IngestionError, match="real directory"):
        discover_inputs(tmp_path / "missing")
    with pytest.raises(IngestionError, match="No CSV/XLSX"):
        load_batch(tmp_path, config_factory())


def test_mixed_inputs_preserve_ids_types_hashes_and_inputs(
    tmp_path, write_csv, write_xlsx, profile_factory, config_factory
):
    csv_path = write_csv(tmp_path / "a.csv")
    xlsx_path = write_xlsx(tmp_path / "z.xlsx")
    before = {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in (csv_path, xlsx_path)}
    result = load_batch(tmp_path, config_factory(profile_factory(), profile_factory("xlsx")))
    assert len(result.accepted) == 2
    assert result.rejected == ()
    assert [row.source.source_file for row in result.rows] == ["a.csv", "z.xlsx"]
    assert [row.source.source_sheet for row in result.rows] == ["", "Sales"]
    for row in result.rows:
        assert row.source.source_row == 2
        assert row.values["branch_id"] == "001"
        assert row.values["transaction_id"] == "000042"
        assert row.values["customer_id"] == "000007"
        assert row.values["amount"] == Decimal("12.34")
        assert row.amount_minor == 1234
    for path, (content, mtime) in before.items():
        assert path.read_bytes() == content
        assert path.stat().st_mtime_ns == mtime
    assert result.files[0].sha256 == hashlib.sha256(before[csv_path][0]).hexdigest()
    frame = result.to_frame()
    assert frame["transaction_id"].tolist() == ["000042", "000042"]
    assert frame["amount"].tolist() == [Decimal("12.34")] * 2
    assert frame["amount_minor"].dtype == object


def test_explicit_mapping_locale_and_normalization(
    tmp_path, write_csv, profile_factory, config_factory
):
    headers = ["Branch", "Receipt", "Day", "Customer", "Total", "Currency", "Note"]
    write_csv(
        tmp_path / "sales.csv",
        headers=headers,
        delimiter=";",
        encoding="cp1250",
        rows=[[" 001 ", "000042", "29/02/2024", "NA", "1.234,50", " pln ", "Łódź"]],
    )
    profile = profile_factory(
        columns=dict(zip(HEADERS, headers[:6], strict=True)),
        encoding="cp1250",
        delimiter=";",
        date_formats=("%d/%m/%Y",),
        decimal_separator=",",
        grouping_separator=".",
        trim_whitespace=True,
        uppercase_currency=True,
        ignored_headers=("Note",),
    )
    row = load_batch(tmp_path, config_factory(profile)).accepted[0]
    assert row.values["branch_id"] == "001"
    assert row.raw_values["Branch"] == " 001 "
    assert row.values["customer_id"] == "NA"
    assert row.raw_values["Note"] == "Łódź"
    assert row.values["currency"] == "PLN"
    assert row.amount_minor == 123450


def test_csv_physical_rows_preamble_blank_and_multiline(tmp_path, profile_factory, config_factory):
    (tmp_path / "sales.csv").write_bytes(
        (
            "preamble\r\n" + ",".join(HEADERS) + "\r\n"
            '001,000042,2024-02-29,"customer\r\ncontinued",12.34,USD\r\n'
            "\r\n001,000043,2024-02-29,000007,1.00,USD\r\n"
        ).encode()
    )
    result = load_batch(tmp_path, config_factory(profile_factory(header_row=2)))
    assert [row.source.source_row for row in result.rows] == [3, 5, 6]
    assert result.rows[0].values["customer_id"] == "customer\r\ncontinued"
    assert [row.status for row in result.rows] == ["accepted", "blank", "accepted"]


@pytest.mark.parametrize(
    "text,match",
    [
        ("", "header row is missing"),
        (",".join(HEADERS) + "\n001,002\n", "row 2"),
        (",".join(HEADERS) + '\n"unterminated', "parsing failed"),
        ("branch_id," + ",".join(HEADERS[1:-1]) + ",branch_id\n", "Duplicate"),
        (",".join(HEADERS[:-1]) + "\n", "Header mismatch"),
        (",".join(HEADERS) + ",extra\n", "unexpected"),
    ],
)
def test_structural_csv_errors_fail_batch(tmp_path, config_factory, text, match):
    (tmp_path / "sales.csv").write_text(text, encoding="utf-8")
    with pytest.raises(IngestionError, match=match) as error:
        load_batch(tmp_path, config_factory())
    assert "sales.csv" in str(error.value)


def test_header_cannot_start_inside_csv_record(tmp_path, profile_factory, config_factory):
    (tmp_path / "sales.csv").write_text('"preamble\ncontinued"\n')
    with pytest.raises(IngestionError, match="inside a multiline"):
        load_batch(tmp_path, config_factory(profile_factory(header_row=2)))


def test_wrong_encoding_does_not_replace_characters(tmp_path, write_csv, config_factory):
    write_csv(tmp_path / "sales.csv", rows=[["Ł", *VALID[1:]]], encoding="cp1250")
    with pytest.raises(IngestionError, match="decoding/parsing failed"):
        load_batch(tmp_path, config_factory())


def test_unmatched_file_blocks_batch(tmp_path, write_csv, write_xlsx, config_factory):
    write_csv(tmp_path / "a.csv")
    write_xlsx(tmp_path / "z.xlsx")
    with pytest.raises(ConfigurationError, match="z.xlsx"):
        load_batch(tmp_path, config_factory())


def test_xlsx_sheet_selection_order_and_physical_rows(
    tmp_path, write_xlsx, profile_factory, config_factory
):
    path = write_xlsx(
        tmp_path / "sales.xlsx",
        sheets={
            "Z": [VALID],
            "A": [[], VALID],
            "Excluded": [["bad"]],
        },
    )
    workbook = load_workbook(path)
    for name in ("A", "Z"):
        workbook[name].insert_rows(1)
        workbook[name]["A1"] = "preamble"
    workbook.save(path)
    workbook.close()
    result = load_batch(
        tmp_path,
        config_factory(
            profile_factory(
                "xlsx",
                sheets=("Z", "A"),
                header_row=2,
            )
        ),
    )
    assert [
        (row.source.source_sheet, row.source.source_row, row.status) for row in result.rows
    ] == [
        ("A", 3, "blank"),
        ("A", 4, "accepted"),
        ("Z", 3, "accepted"),
    ]


def test_numeric_excel_id_rejected_even_with_zero_format(
    tmp_path, write_xlsx, profile_factory, config_factory
):
    path = write_xlsx(tmp_path / "sales.xlsx", rows=[[1, *VALID[1:]]])
    workbook = load_workbook(path)
    workbook["Sales"]["A2"].number_format = "000"
    workbook.save(path)
    workbook.close()
    row = load_batch(tmp_path, config_factory(profile_factory("xlsx"))).rejected[0]
    assert [(issue.code, issue.field) for issue in row.issues] == [("ID_NOT_TEXT", "branch_id")]
    assert row.raw_values["branch_id"] == "1"
    assert row.raw_cells["branch_id"].kind == "numeric"
    assert row.values["branch_id"] is None


def test_formulas_errors_and_boolean_types_are_explicit_rejections(
    tmp_path, write_xlsx, profile_factory, config_factory
):
    write_xlsx(
        tmp_path / "sales.xlsx",
        rows=[
            ["=1+1", VALID[1], VALID[2], True, "#VALUE!", "USD"],
        ],
    )
    row = load_batch(tmp_path, config_factory(profile_factory("xlsx"))).rejected[0]
    assert {(issue.code, issue.field) for issue in row.issues} == {
        ("FORMULA_NOT_ALLOWED", "branch_id"),
        ("ID_NOT_TEXT", "customer_id"),
        ("EXCEL_CELL_ERROR", "amount"),
    }


def test_xlsx_money_uses_stored_numeric_lexeme_not_float(
    tmp_path, write_xlsx, rewrite_xlsx, profile_factory, config_factory
):
    path = write_xlsx(tmp_path / "sales.xlsx", rows=[[*VALID[:4], 1, "USD"]])
    rewrite_xlsx(
        path,
        "xl/worksheets/sheet1.xml",
        b'<c r="E2" t="n"><v>1</v>',
        b'<c r="E2" t="n"><v>123456789012345.67</v>',
    )
    row = load_batch(tmp_path, config_factory(profile_factory("xlsx"))).accepted[0]
    assert row.values["amount"] == Decimal("123456789012345.67")
    assert row.amount_minor == 12345678901234567
    assert row.raw_values["amount"] == "123456789012345.67"


def test_native_excel_date_requires_opt_in_and_rejects_time(
    tmp_path, write_xlsx, profile_factory, config_factory
):
    write_xlsx(
        tmp_path / "sales.xlsx",
        rows=[
            [*VALID[:2], datetime(2024, 2, 29), *VALID[3:]],
            [*VALID[:2], datetime(2024, 2, 29, 12), *VALID[3:]],
        ],
    )
    denied = load_batch(tmp_path, config_factory(profile_factory("xlsx")))
    assert len(denied.rejected) == 2
    allowed = load_batch(tmp_path, config_factory(profile_factory("xlsx", allow_excel_dates=True)))
    assert len(allowed.accepted) == len(allowed.rejected) == 1


def test_missing_sheet_and_corrupt_workbook_fail_batch(
    tmp_path, write_xlsx, profile_factory, config_factory
):
    path = write_xlsx(tmp_path / "sales.xlsx")
    with pytest.raises(IngestionError, match="missing"):
        load_batch(tmp_path, config_factory(profile_factory("xlsx", sheets=("Other",))))
    path.write_bytes(b"not an xlsx")
    with pytest.raises(IngestionError, match="sales.xlsx"):
        load_batch(tmp_path, config_factory(profile_factory("xlsx")))


def test_header_only_source_has_explicit_empty_frame(tmp_path, write_csv, config_factory):
    write_csv(tmp_path / "sales.csv", rows=[])
    result = load_batch(tmp_path, config_factory())
    assert len(result.files) == 1
    assert result.rows == ()
    assert result.to_frame().empty
    assert list(result.to_frame().columns) == [
        *HEADERS,
        "source_file",
        "source_sheet",
        "source_row",
        "amount_minor",
    ]


def test_repeat_load_is_deterministic_and_does_not_deduplicate(tmp_path, write_csv, config_factory):
    write_csv(tmp_path / "sales.csv", rows=[VALID, VALID])
    first = load_batch(tmp_path, config_factory())
    second = load_batch(tmp_path, config_factory())
    assert first == second
    assert len(first.accepted) == 2
    assert len(first.rows) == len(first.accepted) + len(first.rejected) + len(first.blank)


def test_discovery_does_not_hide_permission_errors(tmp_path, monkeypatch):
    def denied_walk(root, onerror, **kwargs):
        onerror(PermissionError("inaccessible branch"))
        return iter(())

    monkeypatch.setattr("sheetmerge.ingest.os.walk", denied_walk)
    with pytest.raises(IngestionError, match="inaccessible branch"):
        discover_inputs(tmp_path)


def test_discovery_rejects_links(tmp_path, monkeypatch):
    linked = tmp_path / "linked.csv"
    linked.touch()
    path_type = type(tmp_path)
    original = path_type.is_symlink
    monkeypatch.setattr(path_type, "is_symlink", lambda path: path == linked or original(path))
    with pytest.raises(IngestionError, match="Links"):
        discover_inputs(tmp_path)


@pytest.mark.parametrize(
    "serial,mac_epoch,accepted",
    [
        (60, False, False),
        (60, True, True),
        (61, False, True),
    ],
)
def test_excel_leap_date_and_epoch(
    tmp_path, write_xlsx, profile_factory, config_factory, serial, mac_epoch, accepted
):
    path = write_xlsx(tmp_path / "sales.xlsx", rows=[[*VALID[:2], serial, *VALID[3:]]])
    workbook = load_workbook(path)
    if mac_epoch:
        workbook.epoch = CALENDAR_MAC_1904
    workbook["Sales"]["C2"].number_format = "yyyy-mm-dd"
    workbook.save(path)
    workbook.close()
    row = load_batch(
        tmp_path, config_factory(profile_factory("xlsx", allow_excel_dates=True))
    ).rows[0]
    assert (row.status == "accepted") is accepted
    if not accepted:
        assert row.issues[0].code == "INVALID_DATE"
        assert "1900-02-29" in row.issues[0].message


def test_excel_date_cannot_hide_fractional_time_through_rounding(
    tmp_path, write_xlsx, rewrite_xlsx, profile_factory, config_factory
):
    path = write_xlsx(tmp_path / "sales.xlsx", rows=[[*VALID[:2], 61, *VALID[3:]]])
    workbook = load_workbook(path)
    workbook["Sales"]["C2"].number_format = "yyyy-mm-dd"
    workbook.save(path)
    workbook.close()
    rewrite_xlsx(path, "xl/worksheets/sheet1.xml", b"<v>61</v>", b"<v>61.000000000001</v>")
    row = load_batch(
        tmp_path, config_factory(profile_factory("xlsx", allow_excel_dates=True))
    ).rejected[0]
    assert row.issues[0].code == "INVALID_DATE"
    assert "time" in row.issues[0].message


def test_stale_excel_dimensions_do_not_hide_records(
    tmp_path, write_xlsx, rewrite_xlsx, profile_factory, config_factory
):
    path = write_xlsx(tmp_path / "sales.xlsx")
    rewrite_xlsx(
        path, "xl/worksheets/sheet1.xml", b'<dimension ref="A1:F2"', b'<dimension ref="A1:A1"'
    )
    assert len(load_batch(tmp_path, config_factory(profile_factory("xlsx"))).accepted) == 1


@pytest.mark.parametrize("number,expected", [("1e-2", 1), ("1.23E+2", 12300), ("-5e-2", -5)])
def test_excel_numeric_scientific_notation_is_exact(
    tmp_path, write_xlsx, rewrite_xlsx, profile_factory, config_factory, number, expected
):
    path = write_xlsx(tmp_path / "sales.xlsx", rows=[[*VALID[:4], 1, "USD"]])
    rewrite_xlsx(
        path,
        "xl/worksheets/sheet1.xml",
        b'<c r="E2" t="n"><v>1</v>',
        f'<c r="E2" t="n"><v>{number}</v>'.encode(),
    )
    row = load_batch(tmp_path, config_factory(profile_factory("xlsx"))).accepted[0]
    assert row.amount_minor == expected


def test_excel_data_beyond_header_fails_instead_of_being_lost(
    tmp_path, write_xlsx, profile_factory, config_factory
):
    write_xlsx(tmp_path / "sales.xlsx", rows=[[*VALID, "unmapped data"]])
    with pytest.raises(IngestionError, match="beyond source headers"):
        load_batch(tmp_path, config_factory(profile_factory("xlsx")))


def test_shipped_configuration_loads():
    from pathlib import Path

    from sheetmerge import load_config

    config = load_config(Path(__file__).parents[1] / "config" / "sources.example.toml")
    assert len(config.profiles) == 2
