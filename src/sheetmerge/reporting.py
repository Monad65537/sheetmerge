"""Exact, verified exports published together in a new output directory."""

import csv
import hashlib
import io
import json
import tempfile
from collections.abc import Mapping
from dataclasses import dataclass, fields, is_dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

from .config import Config
from .consolidate import (
    BUSINESS_KEY,
    ClassifiedRow,
    ConsolidatedBatch,
    source_order,
    transaction_key,
)
from .models import CANONICAL_COLUMNS, PROVENANCE_COLUMNS
from .money import format_minor
from .reconciliation import ReconciliationReport, reconcile
from .version import VERSION

MASTER_COLUMNS = (*CANONICAL_COLUMNS, *PROVENANCE_COLUMNS)
EXCLUSION_COLUMNS = (
    *MASTER_COLUMNS,
    "status",
    "reason_codes",
    "reasons",
    "duplicate_of_file",
    "duplicate_of_sheet",
    "duplicate_of_row",
    "raw_values",
)
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_TEXT = 32_767


class ExportError(ValueError):
    """The complete verified output bundle could not be published."""


@dataclass(frozen=True)
class ExportResult:
    output_directory: Path
    reconciliation: ReconciliationReport


def validate_output_location(input_root: str | Path, output: str | Path) -> Path:
    requested = Path(output).absolute()
    if requested.exists() or requested.is_symlink() or requested.is_junction():
        raise ExportError(f"Output path already exists; choose a new directory: {requested}")
    target = requested.resolve()
    root = Path(input_root).resolve()
    if target == root or target.is_relative_to(root):
        raise ExportError("Output directory must be outside the input tree")
    return target


def _json_value(value: object) -> object:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float):
        return repr(value)  # Audit metadata only; the numeric XML text is retained separately.
    if is_dataclass(value) and not isinstance(value, type):
        return {field.name: _json_value(getattr(value, field.name)) for field in fields(value)}
    if isinstance(value, Mapping):
        return {key: _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    if value is None or isinstance(value, (str, int, bool)):
        return value
    raise ExportError(f"Unsupported audit value type: {type(value).__name__}")


def _json(value: object) -> str:
    return json.dumps(_json_value(value), sort_keys=True, ensure_ascii=True, allow_nan=False)


def _row_order(item: ClassifiedRow) -> tuple:
    return transaction_key(item.row) or ("", ""), source_order(item.row.source)


def _canonical_record(item: ClassifiedRow, config: Config) -> list[str]:
    row = item.row
    values = []
    for field in CANONICAL_COLUMNS:
        value = row.values.get(field)
        if field == "amount" and row.amount_minor is not None:
            value = format_minor(
                row.amount_minor, config.currency_precision[row.values["currency"]]
            )
        elif isinstance(value, (datetime, date)):
            value = value.isoformat()
        values.append("" if value is None else str(value))
    return [*values, row.source.source_file, row.source.source_sheet, str(row.source.source_row)]


def _write_csv(path: Path, headers: tuple[str, ...], rows: list[list[str]]) -> None:
    with path.open("x", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n", quoting=csv.QUOTE_ALL)
        writer.writerow(headers)
        writer.writerows(rows)


def _check_excel_records(records: list[list[str]]) -> None:
    if len(records) + 1 > EXCEL_MAX_ROWS:
        raise ExportError("Master exceeds Excel's worksheet row limit; no files were published")
    for row in records:
        for value in row:
            if len(value.encode("utf-16-le")) // 2 > EXCEL_MAX_TEXT:
                raise ExportError("Master contains text exceeding Excel's cell length limit")
            if ILLEGAL_CHARACTERS_RE.search(value) or "\ufffe" in value or "\uffff" in value:
                raise ExportError("Master contains characters unsupported by Excel/XML")


def _write_workbook(path: Path, records: list[list[str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Master"
    sheet.freeze_panes = "A2"
    all_records = [list(MASTER_COLUMNS), *records]
    for row_number, record in enumerate(all_records, 1):
        for column, value in enumerate(record, 1):
            cell = sheet.cell(row_number, column, value)
            cell.data_type = "s"  # Also forces formula-like identifiers to remain literal text.
            cell.number_format = "@"
            if row_number == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="24476B")
    for column, width in zip("ABCDEFGHI", (24, 24, 18, 24, 28, 12, 44, 24, 14), strict=True):
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = sheet.dimensions
    snapshot = io.BytesIO()
    try:
        workbook.save(snapshot)
    finally:
        workbook.close()
    # Normalize both core and ZIP timestamps. Manifests and exports are reproducible
    # for the same inputs/configuration under the same runtime/library versions.
    with (
        ZipFile(io.BytesIO(snapshot.getvalue())) as source,
        ZipFile(path, "x", compression=ZIP_DEFLATED, compresslevel=9) as target,
    ):
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                # Some Windows XML writers translate LF to CRLF even inside text.
                # Restore the original strings and serialize directly to bytes,
                # encoding CR as a character reference to avoid XML normalization.
                main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                root = ET.fromstring(data)
                for row in root.findall(f"{{{main}}}sheetData/{{{main}}}row"):
                    values = all_records[int(row.attrib["r"]) - 1]
                    for cell in row:
                        text = cell.find(f"{{{main}}}is/{{{main}}}t")
                        if text is not None:
                            column, _ = coordinate_from_string(cell.attrib["r"])
                            text.text = values[column_index_from_string(column) - 1]
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                data = data.replace(b"\r", b"&#13;")
            elif name == "docProps/core.xml":
                root = ET.fromstring(data)
                for field in ("created", "modified"):
                    element = root.find(f"{{http://purl.org/dc/terms/}}{field}")
                    if element is not None:
                        element.text = "2000-01-01T00:00:00Z"
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, data, compresslevel=9)


def _verify_master(stage: Path, records: list[list[str]]) -> None:
    expected = [list(MASTER_COLUMNS), *records]
    with (stage / "master.csv").open(encoding="utf-8", newline="") as stream:
        if list(csv.reader(stream)) != expected:
            raise ExportError("CSV read-back verification failed")
    workbook = load_workbook(stage / "master.xlsx", read_only=True, data_only=False)
    try:
        actual = []
        for cells in workbook["Master"].iter_rows():
            if any(cell.value is not None and cell.data_type != "s" for cell in cells):
                raise ExportError("Excel contains a formula or an inferred nontext value")
            actual.append(["" if cell.value is None else cell.value for cell in cells])
        if actual != expected:
            raise ExportError("Excel read-back verification failed")
    finally:
        workbook.close()


def _write_reports(
    stage: Path, result: ConsolidatedBatch, report: ReconciliationReport, config: Config
) -> None:
    exclusions = []
    with (stage / "audit.jsonl").open("x", encoding="utf-8", newline="\n") as audit:
        for item in sorted(result.rows, key=_row_order):
            audit.write(
                _json(
                    {
                        "source": item.row.source,
                        "status": item.status,
                        "values": item.row.values,
                        "amount_minor": (
                            None if item.row.amount_minor is None else str(item.row.amount_minor)
                        ),
                        "issues": item.issues,
                        "duplicate_of": item.duplicate_of,
                        "raw_values": item.row.raw_values,
                        "raw_cells": item.row.raw_cells,
                    }
                )
                + "\n"
            )
            if item.status == "accepted":
                continue
            winner = item.duplicate_of
            exclusions.append(
                [
                    *_canonical_record(item, config),
                    item.status,
                    _json([issue.code for issue in item.issues]),
                    _json(item.issues),
                    winner.source_file if winner else "",
                    winner.source_sheet if winner else "",
                    str(winner.source_row) if winner else "",
                    _json(item.row.raw_values),
                ]
            )
    _write_csv(stage / "exclusions.csv", EXCLUSION_COLUMNS, exclusions)
    (stage / "reconciliation.json").write_text(
        _json(report.to_dict()) + "\n", encoding="utf-8", newline="\n"
    )
    labels = ("input", "accepted", "rejected", "redundant")
    headers = (
        "currency",
        *(f"{label}_count" for label in labels),
        *(f"{label}_known_amount" for label in labels),
        *(f"{label}_unquantified_count" for label in labels),
        "known_amounts_balanced",
        "monetary_complete",
    )
    _write_csv(
        stage / "reconciliation.csv",
        headers,
        [
            [
                currency.currency,
                *map(str, currency.counts),
                *(format_minor(total, currency.precision) for total in currency.totals),
                *map(str, currency.unquantified),
                "true",
                "true" if currency.unquantified[0] == 0 else "false",
            ]
            for currency in report.currencies
        ],
    )


def export_batch(result: ConsolidatedBatch, config: Config, output: str | Path) -> ExportResult:
    """Reconcile, stage, read back masters, then publish all artifacts together."""
    report = reconcile(result, config)  # No disk writes when conservation fails.
    target = validate_output_location(result.batch.input_root, output)
    records = [_canonical_record(item, config) for item in sorted(result.accepted, key=_row_order)]
    try:
        _check_excel_records(records)
        target.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=f".{target.name}-", dir=target.parent) as staging:
            stage = Path(staging).resolve()
            if stage.parent != target.parent:
                raise ExportError("Staging directory escaped the selected output parent")
            _write_csv(stage / "master.csv", MASTER_COLUMNS, records)
            _write_workbook(stage / "master.xlsx", records)
            _write_reports(stage, result, report, config)
            _verify_master(stage, records)
            outputs = {
                path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                for path in sorted(stage.iterdir())
            }
            manifest = {
                "tool": "sheetmerge",
                "tool_version": VERSION,
                "manifest_version": 1,
                "business_key": BUSINESS_KEY,
                "policies": {
                    "duplicates": "keep earliest file/sheet/row; report others as redundant",
                    "conflicts": "reject all occurrences of a conflicting key",
                    "invalid_candidates": "reject all occurrences of a key with invalid candidates",
                    "blank_rows": "reject with BLANK_ROW",
                },
                "inputs": sorted(result.batch.files, key=lambda item: item.source_file),
                "configuration": config,
                "configuration_sha256": result.batch.config_fingerprint,
                "reconciliation": report.to_dict(),
                "output_sha256": outputs,
            }
            (stage / "manifest.json").write_text(
                _json(manifest) + "\n", encoding="utf-8", newline="\n"
            )
            if target.exists() or target.is_symlink():
                raise ExportError(f"Output path appeared during processing: {target}")
            stage.rename(target)
    except ExportError:
        raise
    except Exception as exc:
        raise ExportError(f"Cannot publish output bundle to {target}: {exc}") from exc
    return ExportResult(target, report)
