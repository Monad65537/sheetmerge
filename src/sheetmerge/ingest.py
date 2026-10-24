"""Read-only, deterministic ingestion with physical row provenance."""

import csv
import hashlib
import io
import os
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import CALENDAR_WINDOWS_1900

from .config import Config, InputLimits, SourceProfile
from .models import BatchResult, CellValue, InputFile, RowResult, SourceRef
from .validation import validate_row
from .xlsx import StoredCell, validate_xlsx


class IngestionError(ValueError):
    """A structural input error prevents trustworthy batch ingestion."""


def discover_inputs(input_root: str | Path) -> tuple[Path, ...]:
    """Discover CSV/XLSX recursively; reject links instead of escaping the root."""
    root = Path(input_root)
    if root.is_symlink() or root.is_junction() or not root.is_dir():
        raise IngestionError(f"Input root must be a real directory: {root}")
    paths = []

    def fail_walk(error: OSError):
        raise error

    try:
        # pathlib globbing can suppress some directory-access errors. An explicit
        # onerror callback prevents an inaccessible branch from being omitted.
        for directory, folders, filenames in os.walk(root, onerror=fail_walk, followlinks=False):
            for name in folders + filenames:
                path = Path(directory) / name
                if path.is_symlink() or path.is_junction():
                    raise IngestionError(f"Links are not supported in input trees: {path}")
            for name in filenames:
                path = Path(directory) / name
                if path.suffix.lower() in (".csv", ".xlsx"):
                    paths.append(path)
    except OSError as exc:
        raise IngestionError(f"Cannot discover inputs in {root}: {exc}") from exc
    return tuple(sorted(paths, key=lambda path: path.relative_to(root).as_posix()))


def _headers(values: list[object], profile: SourceProfile) -> list[str]:
    if not values or any(not isinstance(value, str) or not value.strip() for value in values):
        raise IngestionError("Header cells must be nonempty text")
    if len(set(values)) != len(values):
        raise IngestionError("Duplicate source headers are not allowed")
    missing = set(profile.columns.values()) - set(values)
    unexpected = set(values) - set(profile.columns.values()) - set(profile.ignored_headers)
    if missing or unexpected:
        raise IngestionError(
            f"Header mismatch: missing={sorted(missing)}, unexpected={sorted(unexpected)}"
        )
    return values


def _csv_rows(
    snapshot: bytes, relative: str, profile: SourceProfile, config: Config, records_before: int = 0
) -> list[RowResult]:
    try:
        text = snapshot.decode(profile.encoding)
        reader = csv.reader(
            io.StringIO(text, newline=""),
            delimiter=profile.delimiter,
            quotechar=profile.quotechar,
            strict=True,
        )
        headers = None
        rows = []
        while True:
            start_line = reader.line_num + 1
            try:
                values = next(reader)
            except StopIteration:
                break
            if start_line < profile.header_row:
                if reader.line_num >= profile.header_row:
                    raise IngestionError("header_row lies inside a multiline CSV record")
                continue
            if headers is None:
                headers = _headers(values, profile)
                continue
            config.limits.check(
                "max_total_records", records_before + len(rows) + 1, f"{relative}, row {start_line}"
            )
            if not values:  # A physical blank line is still an accounted-for record.
                values = [""] * len(headers)
            if len(values) != len(headers):
                raise IngestionError(
                    f"row {start_line}: expected {len(headers)} fields, found {len(values)}"
                )
            cells = {
                header: CellValue(value) for header, value in zip(headers, values, strict=True)
            }
            rows.append(validate_row(cells, SourceRef(relative, "", start_line), profile, config))
        if headers is None:
            raise IngestionError("Configured header row is missing")
        return rows
    except (UnicodeError, LookupError, csv.Error) as exc:
        raise IngestionError(f"CSV decoding/parsing failed: {exc}") from exc


def _excel_cell(cell, stored_cells: dict[str, StoredCell], windows_epoch: bool) -> CellValue:
    value = cell.value
    kind = {"f": "formula", "e": "error", "b": "boolean", "n": "numeric", "d": "date"}.get(
        cell.data_type, "text"
    )
    if value is None:
        return CellValue(None)
    stored = stored_cells[cell.coordinate]
    if stored.data_type == "d" and not stored.has_formula:
        kind = "date"
    numeric_text = stored.value_text if stored.data_type == "n" else None
    original_text = (
        stored.value_text if stored.data_type in ("n", "d") and kind != "formula" else None
    )
    invalid_date = False
    if isinstance(value, (date, datetime)) and windows_epoch and numeric_text is not None:
        invalid_date = Decimal(60) <= Decimal(numeric_text) < Decimal(61)
    return CellValue(value, kind, numeric_text, invalid_date, original_text)


def _xlsx_rows(
    snapshot: bytes, relative: str, profile: SourceProfile, config: Config, records_before: int = 0
) -> list[RowResult]:
    decoded_copy, stored_cells = validate_xlsx(
        snapshot,
        profile.sheets,
        config.limits,
        header_row=profile.header_row,
        records_before=records_before,
    )
    workbook = load_workbook(
        io.BytesIO(decoded_copy), read_only=True, data_only=False, keep_links=False
    )
    try:
        missing = set(profile.sheets) - set(workbook.sheetnames)
        if missing:
            raise IngestionError(f"Configured sheets are missing: {sorted(missing)}")
        rows = []
        for name in sorted(profile.sheets):
            sheet = workbook[name]
            # Do not trust possibly stale worksheet dimension metadata.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(min_row=profile.header_row)
            header_cells = next(iterator, ())
            if any(
                cell.value is not None
                and (
                    stored_cells[name][cell.coordinate].data_type not in ("s", "str", "inlineStr")
                    or stored_cells[name][cell.coordinate].has_formula
                )
                for cell in header_cells
            ):
                raise IngestionError(f"sheet {name!r}: header cells must be nonempty text")
            header_values = [cell.value for cell in header_cells]
            # Formatting-only trailing cells do not constitute extra source columns.
            while header_values and header_values[-1] is None:
                header_values.pop()
            try:
                headers = _headers(header_values, profile)
            except IngestionError as exc:
                raise IngestionError(f"sheet {name!r}, row {profile.header_row}: {exc}") from exc
            windows_epoch = workbook.epoch == CALENDAR_WINDOWS_1900
            for row_number, cells in enumerate(iterator, profile.header_row + 1):
                config.limits.check(
                    "max_total_records",
                    records_before + len(rows) + 1,
                    f"{relative}, sheet {name!r}, row {row_number}",
                )
                if any(cell.value is not None for cell in cells[len(headers) :]):
                    raise IngestionError(
                        f"sheet {name!r}, row {row_number}: data extends beyond source headers"
                    )
                mapped_cells = {
                    header: (
                        _excel_cell(cells[index], stored_cells[name], windows_epoch)
                        if index < len(cells)
                        else CellValue(None)
                    )
                    for index, header in enumerate(headers)
                }
                rows.append(
                    validate_row(
                        mapped_cells, SourceRef(relative, name, row_number), profile, config
                    )
                )
        return rows
    finally:
        workbook.close()


def _read_snapshot(path: Path, limits: InputLimits) -> bytes:
    """Check the opened file and bound the read even if its size changes."""
    with path.open("rb") as stream:
        limits.check("max_file_bytes", os.fstat(stream.fileno()).st_size, str(path))
        snapshot = stream.read(limits.max_file_bytes + 1)
    limits.check("max_file_bytes", len(snapshot), str(path))
    return snapshot


def load_batch(input_root: str | Path, config: Config) -> BatchResult:
    """Load a complete batch or raise on structural errors; never write any file."""
    root = Path(input_root).absolute()
    paths = discover_inputs(root)
    if not paths:
        raise IngestionError(f"No CSV/XLSX inputs found in {root}")
    # Resolve all profiles before reading any source, including files without matches.
    sources = [(path, config.profile_for(path.relative_to(root).as_posix())) for path in paths]
    files = []
    rows = []
    for path, profile in sources:
        relative = path.relative_to(root).as_posix()
        try:
            snapshot = _read_snapshot(path, config.limits)
            loader = _csv_rows if profile.format == "csv" else _xlsx_rows
            rows.extend(loader(snapshot, relative, profile, config, records_before=len(rows)))
            files.append(InputFile(relative, hashlib.sha256(snapshot).hexdigest(), profile.name))
        except Exception as exc:
            # Give structural decoder, ZIP/XML, filesystem and workbook failures a
            # consistent public exception, with the input filename always attached.
            raise IngestionError(f"{relative}: {exc}") from exc
    return BatchResult(root, tuple(files), tuple(rows), config.fingerprint)
